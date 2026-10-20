"""Append new behavioral JSON files to behavioral_table.csv.

Reads participant_beh_record2026.xlsx; every row with added == 0 names a JSON
file (same folder) whose trials are appended to behavioral_table.csv, then the
row's `added` is set to 1. The CSV is backed up (.bak) before being written.

Usage: uv run python beh/scripts/update_table.py [data_dir]
"""
import json
import shutil
import sys
from pathlib import Path

import openpyxl
import pandas as pd

DATA_DIR = Path("/home/sebas/data/manualTest")
TABLE, RECORD = "behavioral_table.csv", "participant_beh_record2026.xlsx"
PART_TYPE = {"HC": 1, "CVD": 2, "PD": 3}  # same codes as loader.py


def json_rows(path: Path, row: pd.Series) -> pd.DataFrame:
    raw = json.loads(path.read_text())
    trials = pd.DataFrame(raw["data"], columns=raw["columns"])
    part = [k for k in PART_TYPE if row[k] == 1]
    assert len(part) == 1, f"{row.FolderFile}: expected exactly one of HC/CVD/PD"
    return pd.DataFrame({
        "SubID": row.SUB_ID,
        "Red": trials["Red"],
        "Green": trials["Green"],
        "RunNumber": trials["TrialNumber"],
        "session": row.Session,
        "PartType": PART_TYPE[part[0]],
        "Date": row.Date,
        "FolderOrg": row.FolderFile,
    })


def main(data_dir: Path = DATA_DIR) -> None:
    record = pd.read_excel(data_dir / RECORD)
    pending = record[record["added"] == 0]
    if pending.empty:
        print("nothing to add")
        return

    table = pd.read_csv(data_dir / TABLE)
    new = pd.concat([json_rows(data_dir / row.FolderFile, row) for _, row in pending.iterrows()])
    shutil.copy(data_dir / TABLE, data_dir / (TABLE + ".bak"))
    pd.concat([table, new[table.columns]]).to_csv(data_dir / TABLE, index=False)

    # set `added` in place with openpyxl so the sheet's formatting is kept
    wb = openpyxl.load_workbook(data_dir / RECORD)
    ws = wb.active
    col = [c.value for c in ws[1]].index("added") + 1
    for idx in pending.index:
        ws.cell(row=idx + 2, column=col, value=1)
    wb.save(data_dir / RECORD)
    print(f"added {len(new)} rows from {len(pending)} file(s): {', '.join(pending.FolderFile)}")


if __name__ == "__main__":
    main(Path(sys.argv[1]) if len(sys.argv) > 1 else DATA_DIR)
